#!/usr/bin/env python3
"""
从 WMS 获取货架库存数据，单次同步到 bins_data.xlsx（供真实模式使用）。

is_sim=true（模拟模式）：不修改 bins_data，日志打印"当前为模拟模式"。
is_sim=false（真实模式）：修改 bins_data，日志打印连接状态、获取数量、同步时间。

由 daily_task.sh 每日 04:00 调用（--once）；退出码 0=成功/模拟跳过，1=失败。
日志输出到 logs/lms_YYYYMMDD.log（与 LMS 服务共用同一日志文件）。
"""

import hashlib
import logging
import logging.handlers
import sys
import requests
import openpyxl
import json
from openpyxl.styles import Font, Alignment
from datetime import datetime
from pathlib import Path

# ============================================================
# 路径设置
# ============================================================
SCRIPT_FILE = Path(__file__).resolve()
SCRIPT_DIR = SCRIPT_FILE.parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent.parent
LOG_DIR = PROJECT_ROOT / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)

WMS_URL = "http://10.16.82.150:8082/web/api/wh/reqWmsStock"
BINS_DATA_FILE = SCRIPT_DIR / "bins_data.xlsx"
TEMP_FILE = SCRIPT_DIR / "bins_data_temp.xlsx"

BINS_COLUMNS = [
    "仓库编码",
    "储区编码",
    "储区编码.1",
    "储位编码",
    "储位名称",
    "储位最大容量（万支）",
    "储位状态：0 停用，1 正常，2 仅移入（禁出），3 仅移出（禁入），4 冻结",
    "数量(万支)",
    "品规代码",
    "品规名称",
    "站点",
]

# ============================================================
# 日志设置（复用 LMS 每日日志 logs/lms_YYYYMMDD.log）
# ============================================================
_logger = None


def _get_logger():
    global _logger
    if _logger is not None:
        return _logger
    _logger = logging.getLogger("lms")
    if not _logger.handlers:
        _logger.setLevel(logging.INFO)
        _logger.propagate = False
        today_str = datetime.now().strftime("%Y%m%d")
        fh = logging.handlers.TimedRotatingFileHandler(
            str(LOG_DIR / f"lms_{today_str}.log"),
            encoding="utf-8", when="midnight", interval=1, backupCount=30
        )
        fh.setFormatter(logging.Formatter("%(asctime)s - %(levelname)s - %(message)s", datefmt="%Y-%m-%d %H:%M:%S"))
        _logger.addHandler(fh)
        ch = logging.StreamHandler()
        ch.setFormatter(logging.Formatter("%(asctime)s - %(levelname)s - %(message)s"))
        _logger.addHandler(ch)
    return _logger


def log(msg: str):
    print(msg)
    _get_logger().info(msg)


# ============================================================
# 校验工具
# ============================================================
def _file_checksum(path: Path) -> str:
    if not path.exists():
        return ""
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _data_checksum(data: list) -> str:
    h = hashlib.md5()
    for item in sorted(data, key=lambda x: x.get("binName", "")):
        h.update(repr(item).encode("utf-8"))
    return h.hexdigest()


# ============================================================
# 核心逻辑
# ============================================================
def get_is_sim() -> bool:
    config_path = PROJECT_ROOT / "config.json"
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        return config.get("is_sim", True)
    except Exception:
        log(f"[WARN] 无法读取 {config_path}，默认模拟模式")
        return True


def fetch_wms_stock() -> list:
    log(f"正在连接 WMS 服务器: {WMS_URL}")
    response = requests.get(WMS_URL, timeout=30)
    response.raise_for_status()
    data = response.json()
    log(f"连接成功，获取到 {len(data)} 条储位信息")
    return data


def wms_logical_state_to_bin_status(state: str) -> str:
    mapping = {"0": "1", "1": "1", "2": "2", "3": "3"}
    return mapping.get(str(state), "1")


def _clean_product_name(name) -> str:
    """品规名统一清洗: 全角括号转半角 + 删除所有空白字符(空格/全角空格/换行等)

    与烟箱信息汇总完整版.xlsx 的干净格式对齐, 避免比对端再做字符级兜底。
    """
    s = str(name).replace("（", "(").replace("）", ")")
    return "".join(s.split())


def wms_to_bins_row(item: dict) -> list:
    bin_name = item.get("binName", "")
    product_code = item.get("productCode", "")
    product_name = item.get("productName", "")
    quantity = item.get("quantity", 0)

    if product_code in ("-", "", None):
        product_code = None
    if product_name in ("-", "", None):
        product_name = "无"
    else:
        product_name = _clean_product_name(product_name)
    if quantity == 0 or quantity == "0" or quantity == "-" or quantity == "":
        quantity = "0"

    bin_status = wms_logical_state_to_bin_status(item.get("logicalState", "0"))

    return [
        "110004", "100301", "零烟区", None, bin_name, "50",
        bin_status, quantity, product_code, product_name, None
    ]


def save_to_xlsx(data: list) -> tuple:
    if not data:
        return 0, False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "储位数据"

    for col_idx, header in enumerate(BINS_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, item in enumerate(data, start=2):
        for col_idx, value in enumerate(wms_to_bins_row(item), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value else 0)
            except (ValueError, TypeError):
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    old_checksum = _file_checksum(BINS_DATA_FILE)
    new_checksum = _data_checksum(data)
    changed = old_checksum != new_checksum

    if TEMP_FILE.exists():
        TEMP_FILE.unlink()
    wb.save(str(TEMP_FILE))
    TEMP_FILE.rename(BINS_DATA_FILE)

    return len(data), changed


def sync_job() -> bool:
    """执行一次 WMS 同步。返回 True=成功或模拟模式跳过，False=失败。"""
    is_sim = get_is_sim()

    if is_sim:
        log("当前为模拟模式，不修改 bins_data.xlsx")
        return True

    log("当前为真实模式，开始同步 WMS 数据...")
    try:
        data = fetch_wms_stock()
        if not data:
            log("[WARN] WMS 返回 0 条数据，保留旧 bins_data.xlsx")
            return False
        count, changed = save_to_xlsx(data)
        changed_str = "发生变化" if changed else "无变化"
        log(f"同步完成：获取 {count} 条，状态为 {changed_str}")
        log(f"同步时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        return True
    except requests.exceptions.ConnectionError:
        log(f"连接 WMS 服务器失败: {WMS_URL}")
    except requests.exceptions.Timeout:
        log(f"连接 WMS 服务器超时: {WMS_URL}")
    except Exception as e:
        log(f"同步失败: {e}")
    return False


if __name__ == "__main__":
    # 常驻调度模式已移除：只做单次同步（--once 参数保留兼容，行为相同）
    sys.exit(0 if sync_job() else 1)
